"""Word documents. Heading styles become Markdown headings so the Markdown chunker applies."""

from __future__ import annotations

from pathlib import Path

import docx as python_docx

from behique.extract import Document, register


@register(".docx")
def docx(path: Path) -> Document:
    doc = python_docx.Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.startswith("Heading"):
            level = style.removeprefix("Heading").strip()
            depth = int(level) if level.isdigit() else 1
            lines.append("#" * min(depth, 6) + " " + text)
        elif style == "Title":
            lines.append("# " + text)
        else:
            lines.append(text)
    return Document("\n\n".join(lines), "markdown")
