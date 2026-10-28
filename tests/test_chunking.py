import docx as python_docx
import pytest

from behique import tokens
from behique.chunk import chunk_document, code, markdown
from behique.chunk.pack import HARD_CAP
from behique.extract import Document, UnsupportedFile, extract

LOREM = (
    "Retrieval augmented generation grounds a model in documents. "
    "Chunks are embedded and stored. Queries find the nearest chunks. "
)


def _check_offsets(doc_text, chunks):
    for c in chunks:
        assert doc_text[c.start_char : c.end_char] == c.text
        assert c.token_count == tokens.count(c.text)
        assert c.token_count <= HARD_CAP


def test_markdown_heading_paths():
    text = "# Guide\nintro\n## Install\nrun uv sync\n```\n# not a heading\n```\n## Use\ncall it\n"
    secs = markdown.split(text)
    assert [s.heading_path for s in secs] == ["Guide", "Guide > Install", "Guide > Use"]
    chunks = chunk_document(Document(text, "markdown"))
    _check_offsets(text, chunks)
    assert {c.heading_path for c in chunks} == {"Guide", "Guide > Install", "Guide > Use"}


def test_long_prose_windows_with_overlap_and_cap():
    text = LOREM * 200
    chunks = chunk_document(Document(text, "prose"), chunk_tokens=100, overlap=0.2)
    assert len(chunks) > 5
    _check_offsets(text, chunks)
    assert all(c.token_count <= 100 for c in chunks)
    # consecutive windows overlap
    for a, b in zip(chunks, chunks[1:], strict=False):
        assert b.start_char < a.end_char
    # and together cover the whole text
    assert chunks[0].start_char == 0 and chunks[-1].end_char == len(text.rstrip())


def test_hard_cap_even_with_huge_budget():
    text = "word " * 3000
    chunks = chunk_document(Document(text, "prose"), chunk_tokens=10_000)
    _check_offsets(text, chunks)
    assert max(c.token_count for c in chunks) <= HARD_CAP


def test_no_whitespace_text_still_capped():
    text = "x" * 20_000
    chunks = chunk_document(Document(text, "prose"), chunk_tokens=400)
    _check_offsets(text, chunks)


PY = '''import os

CONST = 1


def alpha(x):
    return x + 1


class Widget:
    """A widget."""

    def spin(self):
        return "spin"

    def stop(self):
        return "stop"
''' + "\n".join(f"    def m{i}(self):\n        return {i} * 2 + os.getpid()\n" for i in range(60))


def test_python_symbols():
    secs = code.split(PY, "python")
    syms = [s.symbol for s in secs]
    assert "alpha" in syms
    assert "Widget.spin" in syms and "Widget.m59" in syms
    chunks = chunk_document(Document(PY, "code", "python"))
    _check_offsets(PY, chunks)
    assert any(c.symbol == "alpha" for c in chunks)
    assert any(c.symbol == "Widget.stop" for c in chunks)
    # import/const glue is kept
    assert any("import os" in c.text for c in chunks)


def test_small_class_kept_whole():
    src = "class A:\n    def f(self):\n        pass\n"
    assert [s.symbol for s in code.split(src, "python")] == ["A"]


def test_js_and_unicode_offsets():
    src = "// héllo ünïcode\nfunction greet(name) {\n  return `hi ${name} 🚀`;\n}\n"
    secs = code.split(src, "javascript")
    assert any(s.symbol == "greet" for s in secs)
    chunks = chunk_document(Document(src, "code", "javascript"))
    _check_offsets(src, chunks)


def test_unknown_language_falls_back():
    assert len(code.split("stuff", "nope-lang")) == 1


def test_empty_documents():
    for kind in ("prose", "markdown"):
        assert chunk_document(Document("   \n", kind)) == []
    assert chunk_document(Document("", "code", "python")) == []


def test_extract_registry(tmp_path):
    (tmp_path / "a.md").write_text("# T\r\nbody\r\n")
    d = extract(tmp_path / "a.md")
    assert d.kind == "markdown" and "\r" not in d.text
    (tmp_path / "b.py").write_text("x = 1\n")
    assert extract(tmp_path / "b.py").language == "python"
    (tmp_path / "c.bin").write_bytes(b"\0")
    with pytest.raises(UnsupportedFile):
        extract(tmp_path / "c.bin")


def test_docx_headings_become_markdown(tmp_path):
    d = python_docx.Document()
    d.add_heading("Top", level=1)
    d.add_paragraph("first para")
    d.add_heading("Sub", level=2)
    d.add_paragraph("second para")
    path = tmp_path / "x.docx"
    d.save(path)
    doc = extract(path)
    assert doc.text.startswith("# Top")
    chunks = chunk_document(doc)
    assert {c.heading_path for c in chunks} == {"Top", "Top > Sub"}


def test_pdf_extract(tmp_path):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(100, 100)
    path = tmp_path / "blank.pdf"
    with open(path, "wb") as f:
        w.write(f)
    assert extract(path).text == ""
